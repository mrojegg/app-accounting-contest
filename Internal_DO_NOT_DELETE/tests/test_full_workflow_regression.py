from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import time
import unittest

import fitz

ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = Path(os.environ.get("SORTING_TEST_SCRIPTS", ROOT / "Internal_DO_NOT_DELETE/scripts"))
sys.path.insert(0, str(ROOT / "Internal_DO_NOT_DELETE/scripts"))
import salary_excel_processor
from openpyxl import Workbook, load_workbook
from sorting_postprocess_settings import SALARY_CREATE_RULE, COMMISSION_EXCEL_RULE, TAX_EXCEL_RULE


class FullWorkflowRegressionTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="sorting-full-cycle-")
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.source = self.root / "source"
        self.output = self.root / "output"
        self.source.mkdir()
        self.output.mkdir()
        self.archive = self.output / "audit"
        self.rules = self.root / "rules.md"
        self.people = {
            "98": ("ОБРАЗЕЦ", "Иванов Иван Иванович"),
            "99": ("ОБРАЗЕЦ", "Демофамилияагб Петр Демофамилияагбич"),
            "102": ("ПРИМЕР", "Сидоров Сидор Сидорович"),
            "109": ("ПРИМЕР", "Орлов Олег Олегович"),
        }
        blocks = ["# Правила\n\n## 11. Офисы\n"]
        for office, (entity, person) in self.people.items():
            categories = ["Зарплата", "Комиссии", "Ндфл и Взносы"] if office in {"98", "109"} else ["Зарплата"]
            category_rules = "\n".join(f"- `{category}`" for category in categories)
            blocks.append(f"### Офис {office}\n\nЮрлицо:\n- `ООО {entity}`\n\n"
                          f"Разрешенные категории по умолчанию:\n{category_rules}\n\n"
                          f"Сотрудники:\n- {person}\n\nКонтрагенты:\n\n")
        blocks.append("## 13. Общие специальные правила\n" + "\n".join(
            f"- {rule}" for rule in [SALARY_CREATE_RULE, COMMISSION_EXCEL_RULE, TAX_EXCEL_RULE]))
        template = (ROOT / "Internal_DO_NOT_DELETE/project_rules/sorting_rules_clean.md").read_text(encoding="utf-8")
        blocks.append(template.split("## 13. Общие специальные правила", 1)[1])
        self.rules.write_text("\n".join(blocks), encoding="utf-8")
        self.env = os.environ | {"SORTING_RULES_FILE": str(self.rules),
                                 "SORTING_TEXT_CACHE_DIR": str(self.root / "cache"),
                                 "PYTHONDONTWRITEBYTECODE": "1"}

    def run_script(self, script, *args, code=0):
        result = subprocess.run([sys.executable, "-B", str(SCRIPTS / script), *map(str, args)],
                                env=self.env, capture_output=True, text=True, timeout=60)
        self.assertEqual(result.returncode, code, result.stdout + result.stderr)
        return result

    def pdf(self, filename, entity, person, date, amount, purpose="Выплата заработной платы. НДФЛ удержан."):
        path = self.source / filename
        with fitz.open() as document:
            page = document.new_page()
            # MuPDF's embedded fallback covers Cyrillic without system fonts.
            page.insert_font(fontname="fixture", fontbuffer=fitz.Font("cjk").buffer)
            page.insert_text((30, 40), f"ПЛАТЕЖНОЕ ПОРУЧЕНИЕ N 10\n{date}\n"
                             f"ООО {entity}\nСумма {amount}\n"
                             f"Сч. N 40817000000000000001 {person} Вид оп. 01\n"
                             f"Назначение платежа {purpose}", fontname="fixture", fontsize=10)
            document.save(path)
        return path

    @unittest.skipUnless(os.environ.get("SORTING_TEST_OCR"), "requires built macOS Vision helper")
    def test_russian_raster_ocr_reads_employee_date_and_amount(self):
        source = self.pdf("text.pdf", "ОБРАЗЕЦ", "Иванов Иван Иванович", "15.08.2026", "1000-00")
        raster = self.root / "raster.pdf"
        with fitz.open(source) as original, fitz.open() as document:
            page = document.new_page(width=original[0].rect.width, height=original[0].rect.height)
            page.insert_image(page.rect, stream=original[0].get_pixmap(matrix=fitz.Matrix(3, 3)).tobytes("png"))
            document.save(raster)
        start = time.monotonic()
        result = subprocess.run([os.environ["SORTING_TEST_OCR"], str(raster), "1"],
                                capture_output=True, text=True, timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("Иванов Иван Иванович", result.stdout)
        self.assertIn("15.08.2026", result.stdout)
        self.assertIn("1000", result.stdout)
        print(f"Vision OCR raster: {time.monotonic() - start:.3f}s, {len(result.stdout)} chars")

    def pipeline(self):
        return self.run_script("run_sorting_pipeline.py", "--source", self.source,
                               "--output", self.output, "--archive-name", "audit", "--mode", "copy")

    def sheet(self, path):
        book = load_workbook(path, read_only=True, data_only=True)
        rows = [list(row) for row in book.active.values]
        book.close()
        return rows

    def test_both_companies_salary_commissions_taxes_and_repeat(self):
        for index, (office, (entity, person)) in enumerate(self.people.items(), start=1):
            self.pdf(f"pay_{office}_1.pdf", entity, person, "15.08.2026", f"{index * 1000}-00")
            self.pdf(f"pay_{office}_2.pdf", entity, person, "29.08.2026", f"{index * 1000 + 500}-00")
        for office in ["98", "109"]:
            entity = self.people[office][0]
            for index, day in enumerate(["15", "29"], start=1):
                self.pdf(f"commission_{office}_{index}.pdf", entity, "Банк", f"{day}.08.2026",
                         f"{index * 100}-00", "Комиссия банка за перевод денежных средств")
            self.pdf(f"tax_{office}.pdf", entity, "Казначейство России", "28.08.2026", "1000-00", "Уплата НДФЛ")
            self.pdf(f"insurance_{office}.pdf", entity, "Казначейство России", "28.08.2026", "2000-00",
                     "Страховые взносы на обязательное пенсионное страхование")
        digests = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in self.source.iterdir()}
        result = self.pipeline()
        self.assertNotIn("postprocess-deferred", result.stdout)
        self.assertEqual(len(list(self.archive.rglob("*.xlsx"))), 10, result.stdout)
        self.assertEqual(len(list((self.archive / "Исходные_чеки").rglob("*.pdf"))), 12)
        for index, office in enumerate(self.people, start=1):
            directory = self.archive / office / "Зарплата"
            rows = self.sheet(directory / f"Зарплата_{office}.xlsx")
            header = next(row for row in rows if row[0] == "ФИО")
            self.assertEqual(header, ["ФИО", "15-08-2026", "29-08-2026", "Общий итог"])
            person_row = next(row for row in rows if row[0] == self.people[office][1])
            self.assertEqual(person_row[1:], [index * 1000, index * 1000 + 500, index * 2000 + 500])
            with fitz.open(directory / f"Зарплата_{office}_чеки.pdf") as merged:
                self.assertEqual(merged.page_count, 2)
                self.assertIn("15.08.2026", merged[0].get_text())
                self.assertIn("29.08.2026", merged[1].get_text())
            self.assertEqual(len(list(directory.glob("*.pdf"))), 1)
        for office in ["98", "109"]:
            entity = self.people[office][0]
            for kind, total in [("налогов", 1000), ("взносов", 2000)]:
                rows = self.sheet(self.archive / office / "Ндфл и Взносы" / f"Распределение_{kind}_{entity}.xlsx")
                self.assertEqual(next(row[1] for row in rows if row[0] == f"Общий итог {kind}"), total)
                self.assertEqual(next(row[1] for row in rows if row[0] == "Общее количество сотрудников"), 2)
                office_rows = [row for row in rows if str(row[0]) in self.people]
                self.assertEqual(sum(row[3] for row in office_rows), total)
            with fitz.open(self.archive / office / "Комиссии" / f"Комиссии_{office}_чеки.pdf") as merged:
                self.assertEqual(merged.page_count, 2)
                self.assertIn("15.08.2026", merged[0].get_text())
        before = {p.relative_to(self.archive): self.sheet(p) for p in self.archive.rglob("*.xlsx")}
        manual_dir = self.archive / "Требует_ручной_проверки"
        manual_dir.mkdir(exist_ok=True)
        self.run_script("reprocess_manual_review.py", "--archive-root", self.archive)
        after = {p.relative_to(self.archive): self.sheet(p) for p in self.archive.rglob("*.xlsx")}
        self.assertEqual(before, after)
        self.assertTrue((self.output / "audit_windows.zip").exists())
        self.assertEqual(digests, {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in self.source.iterdir()})

    def test_triple_duplicate_and_employee_learning(self):
        for day in ["01", "03", "05"]:
            self.pdf(f"retry_{day}.pdf", "ОБРАЗЕЦ", "Иванов Иван Иванович", f"{day}.08.2026", "1000-00")
        self.pipeline()
        manual_dir = self.archive / "Требует_ручной_проверки"
        files = sorted(manual_dir.glob("*.pdf"))
        self.assertEqual(len(files), 3)
        self.assertEqual(list(self.archive.rglob("*.xlsx")), [])
        for path in files[1:]:
            self.run_script("apply_manual_decision.py", "--archive-root", self.archive,
                            "--file-name", path.name, "--action", "exclude_duplicate")
        self.run_script("apply_manual_decision.py", "--archive-root", self.archive,
                        "--file-name", files[0].name, "--action", "move", "--kind", "salary", "--office", "98",
                        "--entity", "ОБРАЗЕЦ", "--category", "Зарплата", "--label", "Иванов Иван Иванович", "--save-to-rules")
        rows = self.sheet(self.archive / "98/Зарплата/Зарплата_98.xlsx")
        self.assertEqual(next(row[-1] for row in rows if row[0] == "ИТОГО"), 1000)
        self.run_script("reprocess_manual_review.py", "--archive-root", self.archive)
        self.assertEqual(len(list((manual_dir / "Исключенные_повторы").glob("*.pdf"))), 2)
        self.assertEqual(rows, self.sheet(self.archive / "98/Зарплата/Зарплата_98.xlsx"))

    def test_unknown_excel_blocks_until_assigned_and_saves_employee(self):
        source = self.source / "ZK-1 от 15.08.2026.xlsx"
        book = Workbook()
        sheet = book.create_sheet("Выплаты")
        sheet.append(["Список № 1 от 15.08.2026"])
        sheet.append([1, "40817000000000000001", None, None, "Новый Иван Иванович", None, None, None, "1 234,56"])
        book.save(source)
        book.close()
        state = self.root / "state.json"
        self.run_script("salary_excel_processor.py", "analyze", "--files", source, "--state-file", state)
        self.assertEqual(len(json.loads(state.read_text())["unknown"]), 1)
        self.run_script("salary_excel_processor.py", "finalize", "--state-file", state, "--output-dir", self.output, code=2)
        self.assertEqual(list(self.output.iterdir()), [])
        assignments = self.root / "assignments.json"
        assignments.write_text(json.dumps({"Новый Иван Иванович": "99"}), encoding="utf-8")
        self.run_script("salary_excel_processor.py", "finalize", "--state-file", state, "--output-dir", self.output,
                        "--assignments-file", assignments)
        payload = json.loads(state.read_text())
        self.assertEqual(payload["created"][0]["amount"], "1234.56")
        self.assertEqual(payload["created"][0]["employees"], 1)
        self.assertEqual(salary_excel_processor.match_employee("Новый Иван Иванович",
                         salary_excel_processor.load_employee_entries(self.rules)), ("99", "Новый Иван Иванович"))

    def test_manual_unknown_employees_waits_for_all_and_learns_rules(self):
        self.pdf("new_one.pdf", "ОБРАЗЕЦ", "Новиков Алексей Сергеевич", "15.08.2026", "1000-00")
        self.pdf("new_two.pdf", "ОБРАЗЕЦ", "Крылов Максим Павлович", "29.08.2026", "2000-00")
        self.pipeline()
        manual_dir = self.archive / "Требует_ручной_проверки"
        files = list(manual_dir.glob("*.pdf"))
        self.assertEqual(len(files), 2)
        for index, (name, office) in enumerate([("Новиков Алексей Сергеевич", "98"), ("Крылов Максим Павлович", "99")]):
            selected = next(path for path in files if name.split()[0].lower() in path.name)
            renamed = manual_dir / "scan.pdf"
            selected.rename(renamed)
            self.run_script("apply_manual_decision.py", "--archive-root", self.archive,
                            "--file-name", renamed.name, "--action", "move", "--kind", "salary", "--office", office,
                            "--entity", "ОБРАЗЕЦ", "--category", "Зарплата", "--label", name, "--save-to-rules")
            self.assertIn(name, self.rules.read_text(encoding="utf-8"))
            if index == 0:
                self.assertEqual(list(self.archive.rglob("*.xlsx")), [])
        self.assertEqual(len(list(self.archive.rglob("*.xlsx"))), 2)
        self.assertEqual(len(list((self.archive / "Исходные_чеки").rglob("*.pdf"))), 2)


if __name__ == "__main__":
    unittest.main()
