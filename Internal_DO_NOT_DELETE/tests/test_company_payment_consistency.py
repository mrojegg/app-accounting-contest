from __future__ import annotations

from decimal import Decimal
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

from openpyxl import load_workbook
import build_commission_table as commissions
import build_insurance_distribution as insurance
import build_insurance_distribution_generic as generic_insurance
import build_insurance_distribution_demo_b as demo_b_insurance
import build_salary_tables as salary
import build_tax_distribution as taxes
import build_tax_distribution_generic as generic_taxes
import build_tax_distribution_demo_b as demo_b_taxes
import salary_employee_counts
import sort_documents as sd
import sorting_budget_payments as budget
import sorting_naming as naming


class CompanyPaymentConsistencyTests(unittest.TestCase):
    def test_execution_date_bank_layouts(self):
        stamps = [
            "ИСПОЛНЕНО: 01.09.2026", "ИСПОЛНЕНО 01.09.2026",
            "ИСПОЛНЕНО\n1 сентября 2026", "01 сентября 2026\nИСПОЛНЕНО",
            "01.09.2026\nИСПОЛНЕНО", "Списано со сч. плат.\n01.09.2026",
            "01.09.2026\nСписано со сч. плат.",
        ]
        for stamp in stamps:
            with self.subTest(stamp=stamp):
                text = "Платежное поручение 28.08.2026\n" + stamp
                for value in [text, sd.normalize_text(text)]:
                    self.assertEqual(naming.extract_payment_execution_date(value), "01-09-2026")
                    for parser in [commissions.parse_date, insurance.parse_date, taxes.parse_date]:
                        self.assertEqual(parser(value), "01.09.2026")
        self.assertIsNone(naming.extract_payment_execution_date("ИСПОЛНЕНО: 31.02.2026"))

    def test_execution_stamp_wins_over_other_dates(self):
        text = "28.08.2026\nСписано со сч. плат.\n28.08.2026\nИСПОЛНЕНО: 01.09.2026"
        self.assertEqual(naming.extract_payment_date(text), "01-09-2026")

    def test_all_categories_ignore_statement_period_in_source_name(self):
        source = Path("VTB_BankStatement_01.08.2026-01.09.2026_page_123.pdf")
        for entity, office in [("ПРИМЕР", "109"), ("ОБРАЗЕЦ", "98")]:
            for category, purpose in [("Зарплата", "Выплата заработной платы"), ("Аренда", "Оплата аренды"),
                                      ("Комиссии", "Комиссия банка"), ("Налоги", "Уплата НДФЛ"),
                                      ("Взносы", "ЕНП (взносы)")]:
                with self.subTest(entity=entity, category=category):
                    text = (f"ПЛАТЕЖНОЕ ПОРУЧЕНИЕ № 12\nДата 28.08.2026\nООО {entity}\n"
                            f"Получатель Иванов Иван Иванович\nНазначение платежа {purpose}\n"
                            "Сумма 1000-00\nИСПОЛНЕНО: 01.09.2026")
                    with patch.object(sd.sorter_core, "file_text", return_value=text):
                        result = sd.safe_filename(source, office, category, entity)
                    self.assertIn("01-09-2026", result)
                    self.assertNotIn("01-08-2026", result)

    def test_budget_kind_distinguishes_enp_insurance_and_other_payments(self):
        cases = [
            ("ЕНП (взносы) НДС не облагается", "Налоги", "insurance"),
            ("Единый налоговый платеж (взносы)", "Налоги", "insurance"),
            ("ЕНП (НДФЛ) НДС не облагается", "Налоги", "tax"),
            ("ЕНП (НДФЛ и взносы)", "Налоги", "other"),
            ("Единый налоговый платеж (по Требованию)", "Налоги", "other"),
            ("Недоимка по Требованию. Регистрационный номер в СФР", "Прочее", "other"),
            ("ЕНП (НДС по корректировке)", "Налоги", "other"),
            ("ЕНП (НДС)", "Налоги", "other"),
        ]
        for text, category, expected in cases:
            with self.subTest(text=text):
                self.assertEqual(budget.distribution_kind(sd.normalize_text(text), category), expected)
        category, _ = sd.detect_category(sd.normalize_text("Платежное поручение ЕНП (взносы) НДС не облагается"), "")
        self.assertEqual(category, "Взносы")

    def test_both_companies_collect_every_budget_document_once(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            texts = {}
            for i, purpose in enumerate(["Уплата НДФЛ", "ЕНП (взносы)", "Недоимка по Требованию", "ЕНП (НДС по корректировке)"]):
                path = root / f"payment_{i}.pdf"
                path.write_bytes(b"fixture")
                texts[path] = f"Платежное поручение 28.08.2026 Сумма {i + 1}-00 {purpose}\nИСПОЛНЕНО: 01.09.2026"
            with patch.object(sd, "file_text", side_effect=texts.__getitem__):
                for module in [taxes, demo_b_taxes, generic_taxes]:
                    rows = module.collect_tax_rows(root)
                    self.assertEqual(rows, [("01.09.2026", 1.0, "payment_0.pdf")])
                for module in [insurance, demo_b_insurance, generic_insurance]:
                    self.assertEqual(module.collect_rows(root), [("01.09.2026", 2.0, "payment_1.pdf")])
                rows = taxes.collect_tax_rows(root)
                output = root / "tax.xlsx"
                taxes.write_distribution(root, output, rows, lambda _: [("98", 1)])
                book = load_workbook(output, read_only=True, data_only=True)
                self.addCleanup(book.close)
                self.assertEqual(book.sheetnames, ["Налоги", "Без распределения"])
                other_rows = list(book["Без распределения"].values)
                self.assertEqual(other_rows[1:3], [("01.09.2026", 3, "payment_2.pdf"), ("01.09.2026", 4, "payment_3.pdf")])
                self.assertEqual(other_rows[-1][1], 7)

    def test_unreadable_budget_document_fails_instead_of_disappearing(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "unknown.pdf"
            path.write_bytes(b"fixture")
            with patch.object(sd, "file_text", return_value="не распознано"):
                with self.assertRaisesRegex(ValueError, "unknown.pdf"):
                    taxes.collect_tax_rows(path.parent, kind="other")

    def test_salary_month_counts_unique_people_by_execution_and_ignores_excluded(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            texts = {}
            people = {"98": ["Иванов Иван Иванович"], "109": ["Иванов Иван Иванович"]}
            for office in people:
                folder = root / "Исходные_чеки" / office / "Зарплата"
                folder.mkdir(parents=True)
                for i, stamp in enumerate(["ИСПОЛНЕНО: 01.09.2026", "01 сентября 2026\nИСПОЛНЕНО"]):
                    path = folder / f"иванов_иван_иванович_28-08-2026_{1000 + i}-00.pdf"
                    path.write_bytes(b"fixture")
                    texts[path] = f"Платежное поручение 28.08.2026 Иванов Иван Иванович\n{stamp}"
            excluded = root / "Требует_ручной_проверки" / "Исключенные_повторы"
            excluded.mkdir(parents=True)
            (excluded / "демофамилияагб_петр_демофамилияагбич_01-09-2026_1000-00.pdf").write_bytes(b"fixture")
            with patch.object(sd, "file_text", side_effect=texts.__getitem__), patch.object(sd, "OFFICE_PEOPLE", people):
                for office in people:
                    self.assertEqual(salary_employee_counts.salary_employee_count(root, office, {"2026-08"}), 0)
                    self.assertEqual(salary_employee_counts.salary_employee_count(root, office, {"2026-09"}), 1)
                parsed = salary.parse_payment_name(next(iter(texts)))
                self.assertEqual(parsed[1:], ("01-09-2026", Decimal("1000.00")))

    def test_distribution_counts_single_payouts_across_months_for_every_builder(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            entities = {"84": "ОБРАЗЕЦ", "98": "ОБРАЗЕЦ", "109": "ПРИМЕР", "76": "ПРИМЕР"}
            people = {office: ["Иванов Иван Иванович", "Демофамилияагб Петр Демофамилияагбич", "Сидоров Сидор Сидорович", "Лишний Из Правил"] for office in entities}
            for office in entities:
                folder = root / "Исходные_чеки" / office / "Зарплата"
                folder.mkdir(parents=True)
                for name in ["иванов_иван_иванович_10-08-2026_1000-00.pdf",
                             "демофамилияагб_петр_демофамилияагбич_01-09-2026_2000-00.pdf",
                             "сидоров_сидор_сидорович_14-08-2026_3000-00.pdf",
                             "сидоров_сидор_сидорович_01-09-2026_4000-00.pdf"]:
                    (folder / name).touch()
                (root / office / "Ндфл и Взносы").mkdir(parents=True)
            excluded = root / "Требует_ручной_проверки" / "Исключенные_повторы"
            excluded.mkdir(parents=True)
            (excluded / "лишний_из_правил_10-08-2026_1000-00.pdf").touch()
            payments = [("14.08.2026", 100.0, "august.pdf"), ("01.09.2026", 300.0, "september.pdf")]
            with patch.object(sd, "file_text", return_value=""), patch.object(sd, "OFFICE_ENTITIES", entities), patch.object(sd, "OFFICE_PEOPLE", people):
                for office, builders, expected in [
                    ("98", [demo_b_taxes, demo_b_insurance], [("84", 3), ("98", 3)]),
                    ("109", [taxes, insurance], [("76", 3), ("109", 3)]),
                ]:
                    tax_root = root / office / "Ндфл и Взносы"
                    for builder in builders:
                        with self.subTest(office=office, builder=builder.__name__):
                            self.assertEqual(builder.office_employee_rows(tax_root, payments), expected)
                    for builder in [generic_taxes, generic_insurance]:
                        self.assertEqual(builder.office_employee_rows(office, tax_root, payments), expected)
                for title, builder in [("Налоги", taxes), ("Взносы", insurance)]:
                    output = root / f'{title}.xlsx'
                    salary_employee_counts.write_archive_distribution(output, payments,
                        lambda _: taxes.office_employee_rows(root / "109" / "Ндфл и Взносы"), builder.build_sheet, title)
                    book = load_workbook(output, read_only=True, data_only=True)
                    self.assertEqual(book.sheetnames, [title])
                    rows = list(book.active.values)
                    book.close()
                    self.assertEqual(next(row[1] for row in rows if row[0] == "Общее количество сотрудников"), 6)
                    office_start = next(i for i, row in enumerate(rows) if row[0] == "Офис") + 1
                    self.assertEqual([(row[0], row[1], row[3]) for row in rows[office_start:]], [("76", 3, 200), ("109", 3, 200)])


if __name__ == "__main__":
    unittest.main()
