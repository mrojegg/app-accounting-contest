from __future__ import annotations

import argparse
from decimal import Decimal
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import apply_manual_decision as manual
import build_salary_tables as salary
import build_tax_distribution as tax
import build_tax_distribution_demo_b as demo_b_tax
import build_insurance_distribution as insurance
import build_insurance_distribution_demo_b as demo_b_insurance
import salary_employee_counts as counts
import salary_excel_processor as excel
import sorting_extraction as extraction
import sorting_file_collection as collection
import sorting_file_ops as file_ops
import sort_documents as sd
from test_extraction_pipeline_integration import runtime_for


class AuditRegressionTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="sorting-regression-")
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)

    def receipt(self, name):
        path = self.root / "98" / "Зарплата" / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"%PDF-1.4\n")
        return path

    def test_legacy_manual_collision_suffix_preserves_receipt(self):
        path = self.receipt("иванов_иван_иванович_15-08-2026_1000-00_2.pdf")
        with patch.object(salary.sd, "file_text", return_value=""):
            self.assertEqual(salary.parse_payment_name(path), ("иванов иван иванович", "15-08-2026", Decimal("1000.00")))

    def test_filename_collision_preserves_amount_contract(self):
        first = self.receipt("иванов_иван_иванович_15-08-2026_1000-00.pdf")
        second = file_ops.unique_destination(first)
        with patch.object(salary.sd, "file_text", return_value=""):
            self.assertIsNotNone(salary.parse_payment_name(second))

    def test_unreadable_receipt_blocks_all_salary_outputs(self):
        self.receipt("scan.pdf")
        with patch.object(salary.sd, "file_text", return_value=""), \
             patch.object(salary, "load_salary_table_settings", return_value=salary.SalaryTableSettings()), \
             patch.object(sys, "argv", ["salary", "--root", str(self.root)]), \
             patch.object(salary, "write_office_outputs") as write:
            with self.assertRaisesRegex(ValueError, "scan.pdf"):
                salary.main()
            write.assert_not_called()

    def test_arbitrary_filename_can_read_payment_fields(self):
        source = self.receipt("scan.pdf")
        text = "Платежное поручение\n15.08.2026\nСумма 1000,00\nПолучатель Иванов Иван Иванович\nНазначение платежа Выплата заработной платы"
        with patch.object(salary.sd, "file_text", return_value=text):
            parsed = salary.parse_payment_name(source)
        self.assertIsNotNone(parsed)
        self.assertEqual(parsed[1:], ("15-08-2026", Decimal("1000.00")))

    def test_same_employee_name_variants_count_once(self):
        self.receipt("иванов_иван_иванович_15-08-2026_1000-00.pdf")
        self.receipt("иванов_иван_29-08-2026_2000-00.pdf")
        with patch.object(salary.sd, "file_text", return_value=""), \
             patch.object(salary.sd, "OFFICE_PEOPLE", {"98": ["Иванов Иван Иванович"]}):
            self.assertEqual(counts.salary_employee_count(self.root, "98", {"2026-08"}), 1)

    def test_new_company_office_participates(self):
        with patch.object(sd, "OFFICE_ENTITIES", {"109": "ПРИМЕР", "120": "ООО ПРИМЕР", "98": "ОБРАЗЕЦ"}):
            self.assertEqual(counts.company_offices("ПРИМЕР", ["109"]), ["109", "120"])

    def test_rounding_preserves_every_kopeck(self):
        result = counts.allocate_money(100, [("76", 1), ("93", 1), ("109", 1)])
        self.assertEqual(sum(result.values()), Decimal("100.00"))
        self.assertEqual(result["76"], Decimal("33.34"))

    def test_no_salary_evidence_blocks_distribution(self):
        for builder in [tax.build_sheet, insurance.build_sheet]:
            with self.assertRaisesRegex(ValueError, "получателей"):
                builder([("28.08.2026", 1000, "receipt.pdf")], [("98", 0)])

    def test_months_share_one_archive_distribution(self):
        rows = [("28.07.2026", 100, "july.pdf"), ("28.08.2026", 200, "august.pdf")]
        seen = []
        def resolve(payments):
            seen.append({counts.month_key(row[0]) for row in payments})
            return [("98", 2)]
        output = self.root / "tax.xlsx"
        counts.write_archive_distribution(output, rows, resolve, tax.build_sheet, "Налоги")
        self.assertEqual(seen, [{"2026-07", "2026-08"}])
        with zipfile.ZipFile(output) as archive:
            workbook = archive.read("xl/workbook.xml").decode()
            sheet = archive.read("xl/worksheets/sheet1.xml").decode()
        self.assertIn('name="Налоги"', workbook)
        self.assertNotIn("2026-07", workbook)
        self.assertNotIn("2026-08", workbook)
        self.assertIn("july.pdf", sheet)
        self.assertIn("august.pdf", sheet)

    def test_insurance_not_included_twice_for_either_company(self):
        source = self.root / "взносы_28-08-2026_2000-00.pdf"
        source.touch()
        text = "Платежное поручение 28.08.2026 Страховые взносы на обязательное пенсионное страхование Сумма 2000,00"
        with patch.object(sd, "file_text", return_value=text):
            for builder in [tax, demo_b_tax]:
                self.assertEqual(builder.collect_tax_rows(self.root), [])
            for builder in [insurance, demo_b_insurance]:
                self.assertEqual(len(builder.collect_rows(self.root)), 1)

    def book(self, name="ZK-1 от 15.08.2026.xlsx", second_sheet=False, amount=1000):
        from openpyxl import Workbook
        path = self.root / name
        book = Workbook()
        sheet = book.create_sheet("Выплаты") if second_sheet else book.active
        sheet.append(["Список № 1 от 15.08.2026"])
        sheet.append([1, "40817000000000000000", None, None, "Иванов Иван Иванович", None, None, None, amount])
        book.save(path)
        book.close()
        return path

    def rules(self):
        path = self.root / "rules.md"
        path.write_text("# Правила\n\n## 11. Офисы\n\n### Офис 98\n\nСотрудники:\n- Иванов Иван Иванович\n\nКонтрагенты:\n\n### Офис 99\n\nСотрудники:\n- Иванов Иван Иванович\n\nКонтрагенты:\n", encoding="utf-8")
        return path

    def test_excel_second_sheet_is_read(self):
        self.assertEqual(len(excel.parse_salary_file(self.book(second_sheet=True))), 1)

    def test_excel_text_amount_is_read(self):
        self.assertEqual(excel.parse_salary_file(self.book(amount="1 000,15"))[0]["amount"], "1000.15")

    def test_legacy_xls_all_sheets_and_numeric_cells(self):
        import xlwt
        path = self.root / "ZK-5 от 15.08.2026.xls"
        book = xlwt.Workbook()
        book.add_sheet("Титул")
        sheet = book.add_sheet("Выплаты")
        sheet.write(0, 0, "Список № 5 от 15.08.2026")
        for col, value in {0: 1, 1: "40817000000000000001", 4: "Иванов Иван Иванович", 8: 1234.56}.items():
            sheet.write(1, col, value)
        book.save(str(path))
        rows = excel.parse_salary_file(path)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["amount"], "1234.56")
        self.assertEqual(rows[0]["source_sheet"], "Выплаты")

    def test_legacy_xls_skips_numbered_secondary_header(self):
        import xlwt
        path = self.root / "ZK-89 от 02.09.2026.xls"
        book = xlwt.Workbook()
        sheet = book.add_sheet("Лист_1")
        sheet.write(0, 0, "Список № 89 от 02.09.2026")
        for col, value in {0: "1", 1: "2", 4: "3", 8: "4"}.items():
            sheet.write(1, col, value)
        for col, value in {0: 1, 1: "40817000000000000002", 4: "Иванов Иван Иванович", 8: 36258.04}.items():
            sheet.write(2, col, value)
        book.save(str(path))

        rows = excel.parse_salary_file(path)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["employee_source"], "Иванов Иван Иванович")
        self.assertEqual(rows[0]["amount"], "36258.04")

    def test_excel_invalid_calendar_date_blocks_import(self):
        with patch.object(excel, "parse_doc_meta", return_value=("ZK-1", "31.02.2026", "")):
            with self.assertRaisesRegex(ValueError, "строка 2"):
                excel.parse_salary_file(self.book())

    def test_excel_invalid_amount_is_not_silently_skipped(self):
        with self.assertRaisesRegex(ValueError, "строка 2"):
            excel.parse_salary_file(self.book(amount="не прочитано"))

    def test_identical_excel_copy_is_rejected(self):
        original = self.book()
        duplicate = self.root / "copy.xlsx"
        shutil.copyfile(original, duplicate)
        with self.assertRaisesRegex(ValueError, "дважды"):
            excel.analyze_files([original, duplicate], self.rules())

    def test_office_assignment_resolves_conflicting_rules(self):
        rules = self.rules()
        self.assertIsNone(excel.match_employee("Иванов Иван Иванович", excel.load_employee_entries(rules)))
        excel.apply_assignments(rules, {"Иванов Иван Иванович": "98"})
        self.assertEqual(excel.match_employee("Иванов Иван Иванович", excel.load_employee_entries(rules)), ("98", "Иванов Иван Иванович"))

    def test_empty_excel_creates_no_output(self):
        with self.assertRaisesRegex(ValueError, "Нет подтвержденных"):
            excel.build_office_workbooks([], self.root)
        self.assertEqual(list(self.root.iterdir()), [])

    def test_pdf_fast_reader_precedes_custom_parser(self):
        path = self.root / "invoice.pdf"
        path.write_bytes(b"%PDF-1.4\n")
        readable = (
            "Счет на оплату. Поставщик ООО Тест. Интернет услуги за август 2026. "
            "Получатель ООО Тест, ИНН 1234567890. Назначение платежа Оплата услуг связи. "
            "Всего к оплате 1000,00 рублей."
        )
        with patch.object(extraction, "read_pymupdf_text", return_value=readable), \
             patch.object(extraction, "extract_pdf_text_layer", side_effect=AssertionError("custom parser called")):
            self.assertIn("Тест", extraction.read_pdf_text(path, runtime=runtime_for(self.root)))

    def test_pdf_native_ocr_does_not_start_quicklook(self):
        binary = self.root / "ocr"
        source = self.root / "scan.pdf"
        with patch.object(extraction, "available_ocr_binary", return_value=binary), \
             patch.object(extraction, "read_quicklook_ocr_text", side_effect=AssertionError("QuickLook called")), \
             patch.object(extraction.platform_support, "pdf_page_count", return_value=1), \
             patch.object(extraction, "_run_command", return_value="Русский текст") as command:
            self.assertEqual(extraction.read_pdf_native_ocr_text(source, runtime=runtime_for(self.root)), "Русский текст")
            command.assert_called_once_with([str(binary), str(source), "1", "--force-ocr", "--quality-json"], timeout=40)

    def test_invoice_total_prefers_payable_over_subtotal(self):
        self.assertEqual(sd.extract_invoice_amount("Итого: 741 857,14 Сумма НДС: 37 092,86 Всего к оплате: 778 950,00", "invoice.pdf"), "778950-00")

    def test_invoice_supplier_not_edo_operator(self):
        text = 'Счет Поставщик (Исполнитель): ООО "ДЕМОПОДРЯДЧИК", ИНН 123 Покупатель: ООО Пример Документ подписан и передан через оператора ЭДО АО «ПФ «СКБ Контур»'
        self.assertEqual(sd.extract_invoice_counterparty(text, "invoice.pdf"), "демоподрядчик")

    def test_utilities_are_not_rent(self):
        category, _ = sd.detect_category("счет возмещение расходов по оплате коммунальных платежей по договору аренды помещения", "счет")
        self.assertEqual(category, "Коммунальные")

    def test_multipage_invoice_is_not_split(self):
        import fitz
        path = self.root / "invoice.pdf"
        with fitz.open() as document:
            document.new_page().insert_text((30, 30), "Invoice first page")
            document.new_page().insert_text((30, 30), "Invoice continuation")
            document.save(path)
        self.assertFalse(collection.is_receipt_bundle(path))
        self.assertIn(path, collection.PRESERVED_MULTIPAGE_PDFS)


if __name__ == "__main__":
    unittest.main()
