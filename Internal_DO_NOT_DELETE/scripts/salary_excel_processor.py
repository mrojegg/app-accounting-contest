#!/usr/bin/env python3
from __future__ import annotations

import argparse
import collections
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
import json
import hashlib
import os
from pathlib import Path
import re
import struct
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape

from sorting_config import parse_office_sections_from_rules
from sorting_text_utils import normalize_text, person_token_matches
import xlsx_builder

# The application carries these pure-Python readers; development uses the same versions.
for dependency_root in (Path(__file__).resolve().parents[1] / "python", Path(__file__).resolve().parents[2] / "Build" / "python-deps"):
    if dependency_root.is_dir():
        sys.path.insert(0, str(dependency_root))


END_OF_CHAIN = 0xFFFFFFFE
FREE_SECTOR = 0xFFFFFFFF
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


@dataclass(frozen=True)
class SalaryRow:
    office: str
    employee: str
    employee_source: str
    account: str
    amount: Decimal
    date: str
    doc: str
    doc_text: str
    file: str
    source_path: str
    source_row: int


def money(value: Decimal | int | str) -> Decimal:
    return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def rules_file_path() -> Path:
    override = os.environ.get("SORTING_RULES_FILE", "").strip()
    if override:
        return Path(override).expanduser()
    return Path(__file__).resolve().parents[1] / "project_rules" / "sorting_rules.md"


def normalize_person(value: str) -> str:
    return " ".join(re.findall(r"[а-яёa-z]+", normalize_text(value)))


def names_equivalent(actual: str, expected: str) -> bool:
    actual_tokens = re.findall(r"[а-яёa-z]+", normalize_text(actual))
    expected_tokens = re.findall(r"[а-яёa-z]+", normalize_text(expected))
    return len(actual_tokens) == len(expected_tokens) and all(
        person_token_matches(expected_token, actual_token)
        for expected_token, actual_token in zip(expected_tokens, actual_tokens)
    )


def load_employee_entries(rules_file: Path) -> list[dict[str, str]]:
    people_map, *_ = parse_office_sections_from_rules(rules_file.read_text(encoding="utf-8"), normalize_text=normalize_text)
    return [
        {"office": office, "name": name, "norm": normalize_person(name)}
        for office, names in people_map.items()
        for name in names
    ]


def match_employee(employee: str, entries: list[dict[str, str]]) -> tuple[str, str] | None:
    normalized = normalize_person(employee)
    exact = [entry for entry in entries if entry["norm"] == normalized]
    matches = exact or [entry for entry in entries if names_equivalent(employee, entry["name"])]
    offices = sorted({entry["office"] for entry in matches}, key=lambda item: int(item)) if matches else []
    if len(offices) != 1:
        return None
    return offices[0], matches[0]["name"]


def parse_doc_meta(doc_text: str, path: Path) -> tuple[str, str, str]:
    match = re.search(r"№\s*(\d+)\s*от\s*(\d{2}\.\d{2}\.\d{4})", doc_text)
    if match:
        return f"ZK-{match.group(1)}", match.group(2), doc_text
    file_match = re.search(r"ZK-(\d+).*?(\d{2}\.\d{2}\.\d{4})", path.name)
    if file_match:
        return f"ZK-{file_match.group(1)}", file_match.group(2), doc_text
    return path.stem, "", doc_text


def column_index(cell_ref: str) -> int:
    letters = re.match(r"[A-Z]+", cell_ref.upper())
    if not letters:
        return 0
    result = 0
    for char in letters.group(0):
        result = result * 26 + ord(char) - 64
    return result - 1


def row_index(cell_ref: str) -> int:
    digits = re.search(r"\d+", cell_ref)
    return int(digits.group(0)) - 1 if digits else 0


def parse_xlsx(path: Path) -> tuple[dict[tuple[int, int], object], str]:
    with zipfile.ZipFile(path) as archive:
        shared_strings: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in root.findall("m:si", NS):
                pieces = [node.text or "" for node in item.findall(".//m:t", NS)]
                shared_strings.append("".join(pieces))
        sheet_name = "xl/worksheets/sheet1.xml"
        root = ET.fromstring(archive.read(sheet_name))

    cells: dict[tuple[int, int], object] = {}
    for cell in root.findall(".//m:c", NS):
        ref = cell.attrib.get("r", "")
        if not ref:
            continue
        cell_type = cell.attrib.get("t", "")
        value_node = cell.find("m:v", NS)
        value: object | None = None
        if cell_type == "s" and value_node is not None:
            index = int(value_node.text or "0")
            value = shared_strings[index] if index < len(shared_strings) else ""
        elif cell_type == "inlineStr":
            value = "".join(node.text or "" for node in cell.findall(".//m:t", NS))
        elif value_node is not None:
            raw = value_node.text or ""
            try:
                value = Decimal(raw)
            except Exception:
                value = raw
        if value is not None:
            cells[(row_index(ref), column_index(ref))] = value
    return cells, "xlsx"


def u16(data: bytes, offset: int) -> int:
    return struct.unpack_from("<H", data, offset)[0]


def u32(data: bytes, offset: int) -> int:
    return struct.unpack_from("<I", data, offset)[0]


def u64(data: bytes, offset: int) -> int:
    return struct.unpack_from("<Q", data, offset)[0]


class CFB:
    def __init__(self, data: bytes) -> None:
        if data[:8] != bytes.fromhex("d0cf11e0a1b11ae1"):
            raise ValueError("not an OLE/CFB file")
        self.data = data
        self.sector_size = 1 << u16(data, 30)
        self.num_fat = u32(data, 44)
        self.first_dir = u32(data, 48)
        difat = [u32(data, 76 + i * 4) for i in range(109)]
        difat = [item for item in difat if item != FREE_SECTOR]
        self.fat: list[int] = []
        for sid in difat[: self.num_fat]:
            sector = self.sector(sid)
            self.fat.extend(struct.unpack("<" + "I" * (len(sector) // 4), sector))
        self.directories: list[tuple[str, int, int, int]] = []
        directory_stream = self.read_chain(self.first_dir)
        for offset in range(0, len(directory_stream), 128):
            entry = directory_stream[offset : offset + 128]
            if len(entry) < 128:
                break
            name_length = u16(entry, 64)
            name = entry[: name_length - 2].decode("utf-16le", "ignore") if name_length >= 2 else ""
            if name:
                self.directories.append((name, entry[66], u32(entry, 116), u64(entry, 120)))

    def sector(self, sid: int) -> bytes:
        offset = (sid + 1) * self.sector_size
        return self.data[offset : offset + self.sector_size]

    def read_chain(self, start: int) -> bytes:
        output = bytearray()
        sid = start
        seen: set[int] = set()
        while sid not in (END_OF_CHAIN, FREE_SECTOR) and sid < len(self.fat) and sid not in seen:
            seen.add(sid)
            output.extend(self.sector(sid))
            sid = self.fat[sid]
        return bytes(output)

    def read_stream(self, name: str) -> bytes:
        for stream_name, _, start, size in self.directories:
            if stream_name.lower() == name.lower():
                return self.read_chain(start)[:size]
        raise KeyError(name)


def biff_records(workbook: bytes):
    position = 0
    while position + 4 <= len(workbook):
        opcode, length = struct.unpack_from("<HH", workbook, position)
        position += 4
        payload = workbook[position : position + length]
        position += length
        yield opcode, payload


def parse_sst(payload: bytes) -> list[str]:
    position = 8
    strings: list[str] = []
    unique_count = struct.unpack_from("<I", payload, 4)[0]
    for _ in range(unique_count):
        char_count = u16(payload, position)
        position += 2
        flags = payload[position]
        position += 1
        is_utf16 = flags & 1
        has_rich = flags & 8
        has_ext = flags & 4
        rich_count = 0
        ext_size = 0
        if has_rich:
            rich_count = u16(payload, position)
            position += 2
        if has_ext:
            ext_size = u32(payload, position)
            position += 4
        byte_count = char_count * (2 if is_utf16 else 1)
        raw = payload[position : position + byte_count]
        position += byte_count
        strings.append(raw.decode("utf-16le" if is_utf16 else "latin1", "replace").strip())
        if has_rich:
            position += 4 * rich_count
        if has_ext:
            position += ext_size
    return strings


def parse_xls(path: Path) -> tuple[dict[tuple[int, int], object], str]:
    workbook = CFB(path.read_bytes()).read_stream("Workbook")
    shared_strings: list[str] = []
    cells: dict[tuple[int, int], object] = {}
    for opcode, payload in biff_records(workbook):
        if opcode == 0x00FC:
            shared_strings = parse_sst(payload)
        elif opcode == 0x00FD:
            row, col, _, index = struct.unpack_from("<HHHI", payload, 0)
            cells[(row, col)] = shared_strings[index] if index < len(shared_strings) else ""
        elif opcode == 0x0203:
            row, col, _ = struct.unpack_from("<HHH", payload, 0)
            cells[(row, col)] = Decimal(str(struct.unpack_from("<d", payload, 6)[0]))
    return cells, "xls"


def parse_salary_file(path: Path) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for sheet_name, sheet_rows in salary_sheets(path):
        texts = [str(value) for row in sheet_rows[:10] for value in row if value is not None]
        metadata = next((text for text in texts if re.search(r"№.*?от\s*\d{2}\.\d{2}\.\d{4}", text)), "")
        doc_no, date, doc_text = parse_doc_meta(metadata, path)
        for row_index, values in enumerate(sheet_rows):
            values = [*values, *([None] * max(0, 9 - len(values)))]
            index, account, employee, amount = values[0], values[1], values[4], values[8]
            index_text = str(index or "").strip()
            if (
                index_text == "1"
                and str(account or "").strip() == "2"
                and str(employee or "").strip() == "3"
                and str(amount or "").strip() in {"4", "4.0"}
            ):
                continue
            if not re.fullmatch(r"\d+(?:\.0+)?", index_text):
                if isinstance(account, str) and re.fullmatch(r"\d{20}", account.strip()) and employee:
                    raise ValueError(f"{path.name}, {sheet_name}, строка {row_index + 1}: не прочитан номер строки выплаты.")
                continue
            if not account and not employee and amount is None:
                continue
            try:
                parsed_amount = money(str(amount).replace("\xa0", "").replace(" ", "").replace(",", "."))
                if not parsed_amount.is_finite() or parsed_amount < 0:
                    raise ValueError("Недопустимая сумма")
                if not isinstance(account, str) or not account.strip() or not isinstance(employee, str) or len(employee.split()) < 2 or not date:
                    raise ValueError("Не прочитаны счет, ФИО или дата")
                datetime.strptime(date, "%d.%m.%Y")
            except Exception as error:
                raise ValueError(f"{path.name}, {sheet_name}, строка {row_index + 1}: проверьте ФИО, счет, дату и сумму.") from error
            rows.append({
                "employee_source": " ".join(employee.split()),
                "account": account.strip(),
                "amount": str(parsed_amount),
                "date": date,
                "doc": doc_no,
                "doc_text": doc_text,
                "file": path.name,
                "source_path": str(path),
                "source_row": row_index + 1,
                "source_sheet": sheet_name,
            })
    if not rows:
        raise ValueError(f"{path.name}: строки выплат не найдены. Проверьте формат ведомости; таблицы не созданы.")
    return rows


def salary_sheets(path: Path):
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                if sheet.max_row and sheet.max_row > 100_000:
                    raise ValueError(f"{path.name}: слишком много строк на листе {sheet.title}")
                yield sheet.title, list(sheet.iter_rows(max_col=9, values_only=True))
        finally:
            workbook.close()
    elif path.suffix.lower() == ".xls":
        import xlrd
        with xlrd.open_workbook(str(path), on_demand=True) as workbook:
            for sheet in workbook.sheets():
                yield sheet.name, [sheet.row_values(index) for index in range(sheet.nrows)]
    else:
        raise ValueError(f"Неподдерживаемый формат ведомости: {path.name}")


def analyze_files(files: list[Path], rules_file: Path) -> dict[str, object]:
    entries = load_employee_entries(rules_file)
    rows: list[dict[str, object]] = []
    unknown: dict[str, dict[str, object]] = collections.defaultdict(lambda: {"rows": 0, "amount": Decimal("0.00"), "files": set()})
    ambiguous: list[dict[str, object]] = []
    fingerprints: dict[str, Path] = {}

    for path in files:
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        if digest in fingerprints:
            raise ValueError(f"Одна ведомость выбрана дважды: {fingerprints[digest].name} и {path.name}. Оставьте один экземпляр во входных файлах.")
        fingerprints[digest] = path
        for raw_row in parse_salary_file(path):
            employee_source = str(raw_row["employee_source"])
            match = match_employee(employee_source, entries)
            if match is None:
                item = unknown[employee_source]
                item["rows"] = int(item["rows"]) + 1
                item["amount"] = money(Decimal(str(item["amount"])) + Decimal(str(raw_row["amount"])))
                item["files"].add(str(raw_row["file"]))
                rows.append({**raw_row, "office": "", "employee": employee_source, "matched": False})
                continue
            office, employee = match
            rows.append({**raw_row, "office": office, "employee": employee, "matched": True})

    return {
        "files": [str(path) for path in files],
        "source_hashes": {str(path): digest for digest, path in fingerprints.items()},
        "rows": rows,
        "unknown": [
            {
                "employee": employee,
                "rows": info["rows"],
                "amount": str(money(info["amount"])),
                "files": sorted(info["files"]),
            }
            for employee, info in sorted(unknown.items())
        ],
        "ambiguous": ambiguous,
    }


def add_employee_to_rules(content: str, office: str, employee: str) -> str:
    if not office.isdigit() or "\n" in employee or "\r" in employee:
        raise ValueError("Некорректное назначение сотрудника")
    pattern = re.compile(rf"(^### Офис {re.escape(office)}\n.*?)(?=^### Офис \d+|^## |\Z)", re.M | re.S)
    match = pattern.search(content)
    if not match:
        raise ValueError(f"Офис {office} не найден в правилах.")
    block = match.group(1)
    if re.search(rf"^- {re.escape(employee)}\s*$", block, re.M):
        return content
    if "Сотрудники:" not in block:
        insert_at = match.start(1) + len(block)
        return content[:insert_at] + f"\nСотрудники:\n- {employee}\n" + content[insert_at:]
    next_header = re.search(r"\n(?:Контрагенты:|Юрлицо:|Разрешенные категории по умолчанию:|### |## )", block[block.find("Сотрудники:") + len("Сотрудники:") :])
    if next_header:
        employee_section_start = block.find("Сотрудники:") + len("Сотрудники:")
        insert_at = match.start(1) + employee_section_start + next_header.start()
    else:
        insert_at = match.start(1) + len(block)
    return content[:insert_at].rstrip() + f"\n- {employee}\n" + content[insert_at:].lstrip("\n")


def apply_assignments(rules_file: Path, assignments: dict[str, str]) -> None:
    content = rules_file.read_text(encoding="utf-8")
    for employee, office in assignments.items():
        employee = employee.strip()
        office = office.strip()
        if employee and office:
            # An explicit office choice replaces conflicting employee entries, never contractor rules.
            lines = content.splitlines(keepends=True)
            current_office = ""
            in_employees = False
            retained = []
            for line in lines:
                header = re.match(r"^### Офис (\d+)", line)
                if header:
                    current_office = header.group(1)
                    in_employees = False
                elif line.strip() == "Сотрудники:":
                    in_employees = True
                elif line.strip() and not line.lstrip().startswith("- "):
                    in_employees = False
                if in_employees and current_office != office and line.strip().startswith("- ") and normalize_person(line.strip()[2:]) == normalize_person(employee):
                    continue
                retained.append(line)
            content = "".join(retained)
            content = add_employee_to_rules(content, office, employee)
    with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=rules_file.parent, delete=False) as staged:
        staged.write(content)
        temporary = Path(staged.name)
    try:
        temporary.replace(rules_file)
    finally:
        temporary.unlink(missing_ok=True)


def date_sort_key(value: str) -> tuple[int, int, int]:
    match = re.match(r"(\d{2})\.(\d{2})\.(\d{4})", value)
    if not match:
        return (9999, 12, 31)
    day, month, year = map(int, match.groups())
    return (year, month, day)


def column_letter(index: int) -> str:
    result = ""
    current = index
    while current > 0:
        current, remainder = divmod(current - 1, 26)
        result = chr(65 + remainder) + result
    return result


def cell(ref: str, value: str, style: int | None = None, cell_type: str = "inlineStr", formula: str | None = None) -> str:
    style_attr = f' s="{style}"' if style is not None else ""
    if formula is not None:
        return f'<c r="{ref}"{style_attr}><f>{escape(formula)}</f><v>{escape(value)}</v></c>'
    if cell_type == "n":
        return f'<c r="{ref}"{style_attr}><v>{escape(value)}</v></c>'
    return f'<c r="{ref}" t="inlineStr"{style_attr}><is><t>{escape(value)}</t></is></c>'


def build_summary_sheet(office: str, rows: list[SalaryRow]) -> str:
    periods = sorted({(row.date, row.doc) for row in rows}, key=lambda item: (date_sort_key(item[0]), item[1]))
    employees = sorted({row.employee for row in rows}, key=lambda value: normalize_text(value))
    headers = ["Сотрудник", *[f"{doc} {date}" for date, doc in periods], "Общий итог"]
    last_col = column_letter(len(headers))
    xml_rows: list[str] = []
    xml_rows.append(f'<row r="1">{cell("A1", f"Офис {office} - зарплата", style=1)}</row>')
    xml_rows.append(f'<row r="2">{cell("A2", f"{len(rows)} строк выплат, {len(employees)} сотрудников")}</row>')
    xml_rows.append(
        '<row r="4">' + "".join(cell(f"{column_letter(index)}4", header, style=1) for index, header in enumerate(headers, start=1)) + "</row>"
    )
    current_row = 5
    totals_by_col = [Decimal("0.00") for _ in headers[1:]]
    for employee in employees:
        row_cells = [cell(f"A{current_row}", employee)]
        employee_total = Decimal("0.00")
        for col_offset, (date, doc) in enumerate(periods, start=2):
            amount = sum((row.amount for row in rows if row.employee == employee and row.date == date and row.doc == doc), Decimal("0.00"))
            totals_by_col[col_offset - 2] += amount
            employee_total += amount
            row_cells.append(cell(f"{column_letter(col_offset)}{current_row}", "" if amount == 0 else f"{amount:.2f}", style=2 if amount else None, cell_type="n" if amount else "inlineStr"))
        totals_by_col[-1] += employee_total
        row_cells.append(cell(f"{last_col}{current_row}", f"{employee_total:.2f}", style=2, cell_type="n"))
        xml_rows.append(f'<row r="{current_row}">{"".join(row_cells)}</row>')
        current_row += 1
    total_row = current_row
    total_cells = [cell(f"A{total_row}", "ИТОГО", style=1)]
    for index, amount in enumerate(totals_by_col, start=2):
        total_cells.append(cell(f"{column_letter(index)}{total_row}", f"{amount:.2f}", style=3, cell_type="n"))
    xml_rows.append(f'<row r="{total_row}">{"".join(total_cells)}</row>')
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f'<dimension ref="A1:{last_col}{total_row}"/>'
        '<sheetViews><sheetView workbookViewId="0"><pane ySplit="4" topLeftCell="A5" activePane="bottomLeft" state="frozen"/></sheetView></sheetViews>'
        '<cols>'
        '<col min="1" max="1" width="38" customWidth="1"/>'
        f'<col min="2" max="{len(headers)}" width="18" customWidth="1"/>'
        '</cols>'
        '<sheetData>' + "".join(xml_rows) + '</sheetData>'
        '</worksheet>'
    )


def build_details_sheet(rows: list[SalaryRow]) -> str:
    headers = ["Офис", "Сотрудник", "Счет", "Сумма", "Дата", "Ведомость", "Файл", "Строка"]
    xml_rows = [
        '<row r="1">' + "".join(cell(f"{column_letter(index)}1", header, style=1) for index, header in enumerate(headers, start=1)) + "</row>"
    ]
    for current_row, row in enumerate(sorted(rows, key=lambda item: (date_sort_key(item.date), item.doc, normalize_text(item.employee))), start=2):
        values = [
            row.office,
            row.employee,
            row.account,
            f"{row.amount:.2f}",
            row.date,
            row.doc,
            row.file,
            str(row.source_row),
        ]
        xml_rows.append(
            f'<row r="{current_row}">'
            + "".join(
                cell(f"{column_letter(index)}{current_row}", value, style=2 if index == 4 else None, cell_type="n" if index == 4 else "inlineStr")
                for index, value in enumerate(values, start=1)
            )
            + "</row>"
        )
    last_row = len(rows) + 1
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f'<dimension ref="A1:H{last_row}"/>'
        '<sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/></sheetView></sheetViews>'
        '<cols>'
        '<col min="1" max="1" width="10" customWidth="1"/>'
        '<col min="2" max="2" width="38" customWidth="1"/>'
        '<col min="3" max="3" width="24" customWidth="1"/>'
        '<col min="4" max="6" width="16" customWidth="1"/>'
        '<col min="7" max="7" width="48" customWidth="1"/>'
        '<col min="8" max="8" width="10" customWidth="1"/>'
        '</cols>'
        '<sheetData>' + "".join(xml_rows) + '</sheetData>'
        '</worksheet>'
    )


def safe_name(value: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', "_", value)


def unique_output_dir(parent: Path) -> Path:
    base = parent / "Зарплата по офисам"
    if not base.exists():
        return base
    index = 2
    while True:
        candidate = parent / f"Зарплата по офисам {index}"
        if not candidate.exists():
            return candidate
        index += 1


def materialize_rows(raw_rows: list[dict[str, object]]) -> list[SalaryRow]:
    return [
        SalaryRow(
            office=str(row["office"]),
            employee=str(row["employee"]),
            employee_source=str(row["employee_source"]),
            account=str(row["account"]),
            amount=money(str(row["amount"])),
            date=str(row["date"]),
            doc=str(row["doc"]),
            doc_text=str(row["doc_text"]),
            file=str(row["file"]),
            source_path=str(row["source_path"]),
            source_row=int(row["source_row"]),
        )
        for row in raw_rows
        if str(row.get("office", "")).strip()
    ]


def build_office_workbooks(rows: list[SalaryRow], output_parent: Path) -> list[dict[str, object]]:
    if not rows:
        raise ValueError("Нет подтвержденных выплат. Таблицы не созданы.")
    target = unique_output_dir(output_parent)
    target.mkdir(parents=True, exist_ok=False)
    created: list[dict[str, object]] = []
    by_office: dict[str, list[SalaryRow]] = collections.defaultdict(list)
    for row in rows:
        by_office[row.office].append(row)
    for office in sorted(by_office, key=lambda value: int(value)):
        office_rows = by_office[office]
        path = target / safe_name(f"Офис {office} - зарплата.xlsx")
        xlsx_builder.build_multi_sheet_workbook_xlsx(
            path,
            [
                ("Свод", build_summary_sheet(office, office_rows)),
                ("Детализация", build_details_sheet(office_rows)),
            ],
        )
        created.append(
            {
                "office": office,
                "file": str(path),
                "rows": len(office_rows),
                "employees": len({row.employee for row in office_rows}),
                "amount": str(money(sum((row.amount for row in office_rows), Decimal("0.00")))),
            }
        )
    return created


def save_state(path: Path, payload: dict[str, object]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def command_analyze(args: argparse.Namespace) -> int:
    files = [Path(item).expanduser().resolve() for item in args.files if not Path(item).name.startswith("~$")]
    payload = analyze_files(files, rules_file_path())
    save_state(args.state_file, payload)
    print(f"Файлов: {len(files)}")
    print(f"Строк выплат: {len(payload['rows'])}")
    print(f"Неизвестных сотрудников: {len(payload['unknown'])}")
    for item in payload["unknown"]:
        print(f"UNKNOWN\t{item['employee']}\t{item['rows']}\t{item['amount']}")
    return 0


def command_finalize(args: argparse.Namespace) -> int:
    state = json.loads(args.state_file.read_text(encoding="utf-8"))
    if not state.get("rows"):
        raise ValueError("В ведомостях нет распознанных выплат. Повторите проверку файлов.")
    for path, digest in state.get("source_hashes", {}).items():
        if hashlib.sha256(Path(path).read_bytes()).hexdigest() != digest:
            raise ValueError("Входная ведомость изменилась после анализа. Повторите проверку файлов.")
    assignments = json.loads(args.assignments_file.read_text(encoding="utf-8")) if args.assignments_file else {}
    missing = [item["employee"] for item in state.get("unknown", []) if not assignments.get(item["employee"], "").strip()]
    if missing:
        print("Не выбран офис для сотрудников:")
        for employee in missing:
            print(employee)
        return 2
    apply_assignments(rules_file_path(), {str(key): str(value) for key, value in assignments.items()})
    refreshed = analyze_files([Path(item) for item in state["files"]], rules_file_path())
    if refreshed["unknown"]:
        save_state(args.state_file, refreshed)
        print("После сохранения правил остались неизвестные сотрудники.")
        return 3
    rows = materialize_rows(refreshed["rows"])
    created = build_office_workbooks(rows, args.output_dir.expanduser().resolve())
    refreshed["created"] = created
    refreshed["output_dir"] = str(Path(created[0]["file"]).parent) if created else ""
    save_state(args.state_file, refreshed)
    print(f"Создано файлов: {len(created)}")
    for item in created:
        print(f"{item['office']}\t{item['file']}\t{item['amount']}")
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    subparsers = parser.add_subparsers(dest="command", required=True)
    analyze = subparsers.add_parser("analyze")
    analyze.add_argument("--state-file", type=Path, required=True)
    analyze.add_argument("--files", nargs="+", required=True)
    finalize = subparsers.add_parser("finalize")
    finalize.add_argument("--state-file", type=Path, required=True)
    finalize.add_argument("--assignments-file", type=Path, required=True)
    finalize.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args(argv)
    if args.command == "analyze":
        return command_analyze(args)
    if args.command == "finalize":
        return command_finalize(args)
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
